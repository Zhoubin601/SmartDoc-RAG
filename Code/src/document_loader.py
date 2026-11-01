# -*- coding: utf-8 -*-
"""
多格式文档解析与切分模块
支持 PDF、DOCX、Markdown、TXT 解析、异常反馈与递归字符分块
"""
from pathlib import Path
from typing import List, Dict, Any, Tuple
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
import pypdf
import docx


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}


def load_single_file(file_path: Path) -> List[Document]:
    """
    解析单个文件并返回 LangChain Document 列表。
    针对空文件、无文本 PDF、损坏文档与非法格式提供明确的异常反馈。
    """
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"文件未找到: {file_path}")

    # 1. 空文件校验
    if file_path.stat().st_size == 0:
        raise ValueError(f"文件「{file_path.name}」大小为 0 字节，属于空文件！")

    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"不支持的文件格式「{suffix}」，目前仅支持: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")

    docs: List[Document] = []

    # 2. PDF 文件解析
    if suffix == ".pdf":
        try:
            reader = pypdf.PdfReader(str(file_path))
            if reader.is_encrypted:
                raise ValueError(f"PDF 文件「{file_path.name}」受密码保护，无法读取内容！")

            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    docs.append(Document(
                        page_content=text.strip(),
                        metadata={"source": file_path.name, "page": i + 1}
                    ))
        except Exception as e:
            if "密码保护" in str(e):
                raise
            raise ValueError(f"PDF 文件「{file_path.name}」解析异常: {e}") from None

        if not docs:
            raise ValueError(f"PDF 文件「{file_path.name}」未提取到有效文本，可能为纯扫描图片或空白 PDF。")

    # 3. DOCX 文件解析
    elif suffix == ".docx":
        try:
            doc = docx.Document(str(file_path))
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            content = "\n".join(full_text)
            if content.strip():
                docs.append(Document(
                    page_content=content.strip(),
                    metadata={"source": file_path.name, "page": 1}
                ))
        except Exception as e:
            raise ValueError(f"Word (.docx) 文件「{file_path.name}」读取损坏或异常: {e}") from None

        if not docs:
            raise ValueError(f"Word (.docx) 文件「{file_path.name}」内容为空或不含文字段落。")

    # 4. TXT / Markdown 文件解析
    elif suffix in [".txt", ".md", ".markdown"]:
        content = ""
        decoded = False
        for encoding in ["utf-8", "utf-8-sig", "gb18030", "gbk", "gb2312", "utf-16"]:
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    content = f.read()
                decoded = True
                break
            except (UnicodeDecodeError, LookupError):
                continue

        if not decoded:
            raise ValueError(f"文本文件「{file_path.name}」编码格式不受支持，无法正确解码。")

        if content.strip():
            docs.append(Document(
                page_content=content.strip(),
                metadata={"source": file_path.name, "page": 1}
            ))
        else:
            raise ValueError(f"文本文件「{file_path.name}」内容仅包含空白字符。")

    return docs


def split_documents(
    documents: List[Document],
    chunk_size: int = 600,
    chunk_overlap: int = 100
) -> List[Document]:
    """
    对文档集合进行智能递归字符切分
    """
    if chunk_size <= 0 or chunk_overlap < 0 or chunk_overlap >= chunk_size:
        raise ValueError("切片大小必须为正数，重叠跨度必须非负且小于切片大小。")
    if not documents:
        return []

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", "。", "！", "？", "；", " ", ""]
    )
    return splitter.split_documents(documents)


def load_and_split_files(
    file_paths: List[Path],
    chunk_size: int = 600,
    chunk_overlap: int = 100,
    source_names: Dict[str, str] = None,
) -> Tuple[List[Document], Dict[str, Any]]:
    """
    批量加载并切分所有文件，返回切分后的 Document 列表以及明细解析报告
    """
    all_raw_docs: List[Document] = []
    report: Dict[str, Any] = {
        "success": [],
        "failed": []
    }

    for fp in file_paths:
        fp = Path(fp)
        source = (source_names or {}).get(fp.name, fp.name)
        try:
            docs = load_single_file(fp)
            for doc in docs:
                doc.metadata["source"] = source
                if source_names is not None:
                    doc.metadata["storage_id"] = fp.name
            all_raw_docs.extend(docs)
            total_chars = sum(len(d.page_content) for d in docs)
            report["success"].append({
                "file": source,
                "pages": len(docs),
                "chars": total_chars
            })
        except Exception as e:
            report["failed"].append({
                "file": source,
                "reason": str(e)
            })

    split_docs = split_documents(all_raw_docs, chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    return split_docs, report

