# -*- coding: utf-8 -*-
"""
文档加载与切分模块单元测试
"""
import pytest
from pathlib import Path
from langchain_core.documents import Document
from src.document_loader import (
    load_single_file,
    split_documents,
    load_and_split_files,
    SUPPORTED_EXTENSIONS
)


def test_load_txt_file(tmp_path: Path):
    file_path = tmp_path / "sample.txt"
    file_path.write_text("这是开源软件构建技术的测试文本内容。\n包含第二行说明。", encoding="utf-8")

    docs = load_single_file(file_path)
    assert len(docs) == 1
    assert "开源软件构建技术" in docs[0].page_content
    assert docs[0].metadata["source"] == "sample.txt"
    assert docs[0].metadata["page"] == 1


def test_load_md_file(tmp_path: Path):
    file_path = tmp_path / "sample.md"
    file_path.write_text("# 标题一\n这是 Markdown 测试段落。\n\n## 标题二\n第二段内容。", encoding="utf-8")

    docs = load_single_file(file_path)
    assert len(docs) == 1
    assert "Markdown 测试段落" in docs[0].page_content
    assert docs[0].metadata["source"] == "sample.md"


def test_load_empty_file_raises(tmp_path: Path):
    file_path = tmp_path / "empty.txt"
    file_path.write_text("", encoding="utf-8")

    with pytest.raises(ValueError, match="大小为 0 字节"):
        load_single_file(file_path)


def test_load_whitespace_only_file_raises(tmp_path: Path):
    file_path = tmp_path / "blank.txt"
    file_path.write_text("   \n\n\t  ", encoding="utf-8")

    with pytest.raises(ValueError, match="仅包含空白字符"):
        load_single_file(file_path)


def test_load_unsupported_format_raises(tmp_path: Path):
    file_path = tmp_path / "image.png"
    file_path.write_bytes(b"\x89PNG\r\n\x1a\n")

    with pytest.raises(ValueError, match="不支持的文件格式"):
        load_single_file(file_path)


def test_split_documents():
    long_text = "段落一内容。" * 50 + "\n\n" + "段落二内容。" * 50
    doc = Document(page_content=long_text, metadata={"source": "long.txt"})

    chunks = split_documents([doc], chunk_size=200, chunk_overlap=30)
    assert len(chunks) > 1
    for chunk in chunks:
        assert len(chunk.page_content) <= 250
        assert chunk.metadata["source"] == "long.txt"


def test_load_and_split_files_batch_report(tmp_path: Path):
    # 正常文件 1
    f1 = tmp_path / "doc1.txt"
    f1.write_text("文档一的核心内容与关键参数说明。", encoding="utf-8")

    # 正常文件 2
    f2 = tmp_path / "doc2.md"
    f2.write_text("# 文档二\n详细架构与实施指南。", encoding="utf-8")

    # 异常空文件
    f3 = tmp_path / "empty.txt"
    f3.write_text("", encoding="utf-8")

    split_docs, report = load_and_split_files([f1, f2, f3], chunk_size=500, chunk_overlap=50)

    assert len(report["success"]) == 2
    assert len(report["failed"]) == 1
    assert report["failed"][0]["file"] == "empty.txt"
    assert "0 字节" in report["failed"][0]["reason"]
    assert len(split_docs) >= 2


def test_real_pdf_text_and_page_metadata(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    for text in ["Port 7219", "Backup 17 days"]:
        page = writer.add_blank_page(width=400, height=400)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 300 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    path = tmp_path / "manual.pdf"
    writer.write(path)
    docs = load_single_file(path)
    assert len(docs) == 2
    assert "7219" in docs[0].page_content and docs[1].metadata["page"] == 2


@pytest.mark.parametrize("encrypted", [False, True])
def test_blank_or_encrypted_pdf(tmp_path, encrypted):
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    if encrypted:
        writer.encrypt("password")
    path = tmp_path / "empty.pdf"
    writer.write(path)
    with pytest.raises(ValueError, match="密码保护" if encrypted else "未提取到有效文本"):
        load_single_file(path)


def test_real_docx_paragraphs(tmp_path):
    from docx import Document as WordDocument
    doc = WordDocument()
    doc.add_paragraph("端口为 7219。")
    doc.add_paragraph("备份周期为 17 天。")
    path = tmp_path / "manual.docx"
    doc.save(path)
    loaded = load_single_file(path)
    assert len(loaded) == 1
    assert "7219" in loaded[0].page_content and "17 天" in loaded[0].page_content
    assert loaded[0].metadata["source"] == "manual.docx"


@pytest.mark.parametrize("suffix", [".pdf", ".docx"])
def test_corrupted_binary_file(tmp_path, suffix):
    path = tmp_path / ("corrupt" + suffix)
    path.write_bytes(b"invalid binary format")
    with pytest.raises(ValueError, match="异常|损坏"):
        load_single_file(path)


@pytest.mark.parametrize("size,overlap", [(100,100), (100,101), (0,0), (100,-1)])
def test_invalid_chunk_parameters(size, overlap):
    with pytest.raises(ValueError, match="切片大小"):
        split_documents([Document(page_content="测试文本")], size, overlap)
